"""Per-paper sections: a divider per research front, then 2-3 figure-led slides per paper.

Slide roles per paper
  method   the paper's overview / architecture / illustration figure + what it does
  results  its main results figure or table + the numbers that matter
  us       (optional) what it means for RSA / maps / canvas, with a second figure if useful

Figures are chosen automatically from the paper's caption candidates (see `pick`),
or pinned by label in paper_content.py once a human has read the paper. Statements
come from paper_content.py; a paper without written content gets its auto-selected
figures and an amber "notes pending" line, never invented text.
"""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from PIL import Image

from . import paper_content as PC
from .papers import REGISTRY
from .pptx_kit import (AMBER, BG2, BODY_PT, CAPTION_PT, CONTENT_TOP, HEAD_FONT, INK, LIGHT_TEXT, MARGIN, MUTED,
                       SLIDE_W, TEAL, add_text, box, set_bg)
from .slides import CW, chrome, figure, finish_footer, new_slide, statements

FRONTS = [
    ("filtered_ann", "Filtered ANN", "Fast vector search under attribute filters: where RSA's predicate plugs in."),
    ("negation", "Negation and logic in retrieval", "Exclusion, negation and instructions: the bar for RSA's log(1 − p) algebra."),
    ("sae", "Sparse autoencoders and interpretable embeddings", "Concept vocabularies from embeddings: RSA predicates and map axes."),
    ("cir", "Composed image retrieval", "Edit-by-text retrieval in fashion: the prior art for the intent canvas."),
    ("genret", "Generative retrieval in e-commerce", "Semantic IDs and generated candidates: maps as codes, masks as pruning."),
    ("production", "Production search, QU and conversation", "How deployed systems structure queries, facets and dialogue."),
]

ROLE_KEYWORDS = {
    "method": ["overview", "framework", "architecture", "pipeline", "illustration", "workflow", "example",
               "our method", "proposed", "diagram", "model", "approach", "taxonomy"],
    "results": ["result", "performance", "recall", "qps", "throughput", "accuracy", "comparison", "compared",
                "ndcg", "mrr", "speedup", "latency", "a/b", "online", "ablation", "benchmark", "selectivity"],
}


def pick(ctx, key: str, role: str, used: set, label: str | None = None):
    """Best unused crop for a role: an explicit label wins, else caption-keyword score."""
    cs = [c for c in ctx.cands if c.paper == key and c.file and c.file not in used]
    if label:
        for c in cs:
            if c.label == label:
                return c
        return None
    if not cs:
        return None

    def score(c):
        cap = c.caption.lower()
        kw = ROLE_KEYWORDS[role]
        sc = sum(len(kw) - i for i, k in enumerate(kw) if k in cap)
        if role == "method":
            sc += 3 if c.kind == "fig" else -5
            sc += 2 if c.number <= 2 else 0
        else:
            sc += 1 if c.kind == "tab" else 0
        return sc

    return max(cs, key=score)


def _aspect(ctx, c) -> float:
    with Image.open(ctx.fig_dir / c.file) as im:
        return im.width / im.height


def _divider(ctx, prs, front_key, name, blurb, papers):
    s = new_slide(prs)
    set_bg(s, BG2)
    add_text(s, MARGIN, 1.1, 10, 0.35, "RESEARCH FRONT", size=13, color=TEAL, bold=True, char_spacing=1.5)
    add_text(s, MARGIN, 1.6, 11.5, 1.4, name, size=40, font=HEAD_FONT, color=INK)
    add_text(s, MARGIN, 3.0, 11.5, 0.8, blurb, size=20, color=MUTED)
    lines = [f"**{p.short}** · {p.authors} {p.year}" for p in papers]
    half = (len(lines) + 1) // 2
    add_text(s, MARGIN, 4.1, 5.9, 2.6, lines[:half], size=BODY_PT, space_after=6)
    add_text(s, MARGIN + 6.2, 4.1, 5.9, 2.6, lines[half:], size=BODY_PT, space_after=6)
    add_text(s, SLIDE_W - MARGIN - 0.6, 6.9, 0.6, 0.3, str(s._n), size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = f"Section: {name}. {blurb} Papers: " + "; ".join(p.short for p in papers)


def _paper_slide(ctx, prs, p, role, spec, cand, sid):
    s = new_slide(prs)
    kicker = f"{p.short} · {p.authors} {p.year}"
    default_title = {"method": "What it does", "results": "What it shows", "us": "What it means for us"}[role]
    title = spec.get("title") or default_title
    footer = f"{p.authors} ({p.year}), arXiv:{p.arxiv}"
    chrome(ctx, s, sid, kicker, title, footer)
    items = spec.get("statements") or [
        "__Notes pending:__ figure auto-selected from the paper's captions; statements are written only after "
        "the PDF has been read (see qa_report.md)."]
    want = spec.get("want") or {"method": "method overview figure", "results": "main results figure or table",
                                "us": "second figure"}[role]
    wide = cand is not None and _aspect(ctx, cand) > 2.1
    if wide:
        # full-width figure, statements underneath
        figure(ctx, s, sid, p.key, MARGIN, CONTENT_TOP, CW, 3.0, want, cand=cand, cap_h=0.3)
        statements(s, MARGIN, 5.05, CW, 1.6, items, gap=6)
    else:
        figure(ctx, s, sid, p.key, MARGIN, CONTENT_TOP, 7.9, 4.45, want, cand=cand, cap_h=0.5)
        statements(s, 8.85, CONTENT_TOP, 3.9, 4.9, items)
    if role == "us" and spec.get("links_to"):
        tx = 8.85 if not wide else MARGIN
        ty = 6.2 if not wide else 4.6
        for i, idea in enumerate(spec["links_to"]):
            box(s, tx + i * 1.3, ty, 1.2, 0.4, idea, fill=TEAL, line=None, color=LIGHT_TEXT, size=12)
    finish_footer(s)
    s.notes_slide.notes_text_frame.text = spec.get("notes") or (
        f"{p.title or p.expected_title} ({p.cite}). Figure: "
        f"{cand.label + ' — ' + cand.caption[:300] if cand else 'placeholder, PDF not available'}.")
    return s


def build_front(ctx, prs, front_key):
    name, blurb = next((n, b) for k, n, b in FRONTS if k == front_key)
    papers = [p for p in REGISTRY if p.front == front_key]
    _divider(ctx, prs, front_key, name, blurb, papers)
    for p in papers:
        spec = PC.CONTENT.get(p.key, {})
        used: set = set()
        roles = spec.get("roles") or ["method", "results"]
        for role in roles:
            rs = spec.get(role, {})
            cand = pick(ctx, p.key, "results" if role == "us" else role, used, rs.get("figure"))
            if cand is not None:
                used.add(cand.file)
            _paper_slide(ctx, prs, p, role, rs, cand, f"{p.key}:{role}")
