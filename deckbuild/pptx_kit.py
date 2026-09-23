"""Small python-pptx toolkit: design tokens, text, panels, native diagram shapes."""
from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---- design tokens -----------------------------------------------------------
BG, BG2 = "F5F2EC", "EAE5DB"
INK, MUTED = "1E2328", "5B636B"
TEAL, AMBER = "1F6F6B", "A5541B"
DARK = "1E2328"
PANEL, BORDER = "FFFDF8", "DDD6C8"
LIGHT_TEXT = "F5F2EC"
HEAD_FONT, BODY_FONT = "Libre Baskerville", "IBM Plex Sans"

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
CONTENT_TOP, CONTENT_BOTTOM = 1.62, 6.62
FOOTER_Y = 6.9
BODY_PT, FOOTER_PT, CAPTION_PT, DIAGRAM_PT = 18, 12, 12, 15


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def set_bg(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def _style_run(run, size, color, font, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic


MARK = re.compile(r"(\*\*.+?\*\*|__.+?__)")


def add_text(slide, x, y, w, h, paras, size=BODY_PT, color=INK, font=BODY_FONT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=0, line_spacing=None,
             italic=False, char_spacing=None, accent=TEAL, link=None, name=None):
    """Text box. `paras` is a str or list of str; **x** = bold accent, __x__ = amber bold."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for part in MARK.split(ptxt):
            if not part:
                continue
            r = p.add_run()
            if part.startswith("**"):
                r.text = part[2:-2]
                _style_run(r, size, accent, font, True, italic)
            elif part.startswith("__"):
                r.text = part[2:-2]
                _style_run(r, size, AMBER, font, True, italic)
            else:
                r.text = part
                _style_run(r, size, color, font, bold, italic)
            if char_spacing is not None:
                r.font._element.set("spc", str(int(char_spacing * 100)))
            if link:
                r.hyperlink.address = link
    return tb


def set_alpha(shape, alpha_pct: int) -> None:
    """Fill transparency (python-pptx has no API): alpha in percent opacity."""
    sp = shape.fill._xPr.find(qn("a:solidFill"))
    clr = sp[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha_pct * 1000)))


def _drop_style(shape) -> None:
    """Remove the theme style reference (kills default shadow/3-D effects)."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(slide, x, y, w, h, text="", fill=PANEL, line=BORDER, color=INK, size=DIAGRAM_PT,
        bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, align=PP_ALIGN.CENTER,
        font=BODY_FONT, line_w=1.0, dash=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    _drop_style(s)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    if text:
        lines = text if isinstance(text, list) else [text]
        for i, t in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            for part in MARK.split(t):
                if not part:
                    continue
                r = p.add_run()
                if part.startswith("**"):
                    r.text = part[2:-2]
                    _style_run(r, size, color, font, True)
                else:
                    r.text = part
                    _style_run(r, size, color, font, bold)
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5, dash=None, head=True, begin=None, end=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if begin is not None:
        c.begin_connect(*begin)
    if end is not None:
        c.end_connect(*end)
    _drop_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if head:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle")
        te.set("w", "med")
        te.set("len", "med")
    return c


def picture_fit(slide, path: Path, x, y, w, h, align="center"):
    """Insert an image scaled to fit (x, y, w, h) without distortion."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph)), (px, py, pw, ph)


def equation(slide, png: Path, x, y, max_w, max_h=None, pt=20, align="left"):
    """Place a 300-dpi equation PNG (rendered at 22 pt) so its glyphs read at ~`pt`."""
    with Image.open(png) as im:
        iw, ih = im.size
    w_in, h_in = iw / 300 * pt / 22, ih / 300 * pt / 22
    s = 1.0
    if w_in > max_w:
        s = max_w / w_in
    if max_h and h_in * s > max_h:
        s = max_h / h_in
    w_in, h_in = w_in * s, h_in * s
    px = x + (max_w - w_in) / 2 if align == "center" else x
    pic = slide.shapes.add_picture(str(png), Inches(px), Inches(y), Inches(w_in), Inches(h_in))
    return pic, (px, y, w_in, h_in, pt * s)
